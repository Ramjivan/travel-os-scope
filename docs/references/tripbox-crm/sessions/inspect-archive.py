from pathlib import Path
import hashlib, json, shutil
from PIL import Image
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
source = Path('C:/Users/ramji/Downloads/tos-audit-fictional-report-b-2026-10-10.xlsx')
dest = root / 'documents' / source.name
dest.parent.mkdir(exist_ok=True)
if not dest.exists():
    shutil.copy2(source, dest)
assert hashlib.sha256(source.read_bytes()).digest() == hashlib.sha256(dest.read_bytes()).digest()
wb = load_workbook(dest, data_only=False)
result = {'source':source.name, 'archive':dest.name, 'bytes':dest.stat().st_size,
          'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(), 'sheets':[]}
for sheet in wb:
    rows = [[cell.value for cell in row] for row in sheet]
    result['sheets'].append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,
        'merged':list(map(str,sheet.merged_cells.ranges)), 'values':rows})
(root / 'documents' / 'report-b-extraction.json').write_text(json.dumps(result,indent=2,default=str),encoding='utf-8')
manifest_file = root / 'screenshots' / 'manifest.json'
manifest = json.loads(manifest_file.read_text(encoding='utf-8'))
known = {m['archive_filename'] for m in manifest}
for file in sorted((root / 'screenshots').glob('tripbox-*.jpg')):
    if file.name in known: continue
    with Image.open(file) as img:
        img.load()
        dimensions = list(img.size)
    number = file.name.split('-')[1]
    manifest.append({'reference_id':'TRIPBOX-'+number,'source':'Authorised live browser audit',
        'date':'2026-10-10','role':'Owner','archive_filename':file.name,
        'archive_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'dimensions':dimensions,
        'redactions':[], 'limitations':'Fictional audit records; full-page images may clip internally scrolling panels. Text evidence records exact values.'})
manifest_file.write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
for item in manifest:
    file=root/'screenshots'/item['archive_filename']
    assert hashlib.sha256(file.read_bytes()).hexdigest()==item['archive_sha256']
    with Image.open(file) as img:
        img.load()
        assert list(img.size)==item['dimensions']
files = sorted((root/'sessions'/'evidence').glob('L*.txt'))
assert len(files)==len({p.name.split('-')[0] for p in files})
index='# TripboxCRM raw browser evidence\n\nOwner, 9–10 October 2026. Captures include loading and unsuccessful drafts; use the linked screen records for committed outcomes. Demo email redacted. Numeric IDs are unique; filenames preserve the inspected state.\n\n| ID | Raw snapshot |\n| --- | --- |\n'
for file in files:
    index+='| '+file.name.split('-')[0]+' | ['+file.stem[5:]+'](evidence/'+file.name+') |\n'
(root/'sessions'/'evidence-index.md').write_text(index,encoding='utf-8')
print(json.dumps(result,indent=2,default=str))
print('Verified',len(manifest),'decoded screenshots;',len(files),'unique raw snapshots')
